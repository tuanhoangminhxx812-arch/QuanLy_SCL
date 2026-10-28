"""
thqt_parser.py
Module đọc, bóc tách và quản lý file Bảng Tổng hợp Quyết toán (THQT) kinh phí SCL (.xlsx)
Phục vụ tự động trích xuất các chỉ tiêu tài chính và điền vào:
- Báo cáo Thẩm tra
- Tờ trình duyệt quyết toán
- Quyết định phê duyệt quyết toán
"""

import os
import re
import json
from io import BytesIO
import openpyxl
import pandas as pd
from form_module import doc_so_vn

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
os.makedirs(DATA_DIR, exist_ok=True)


def _clean_num(val):
    if val is None or val == "":
        return 0
    try:
        if isinstance(val, (int, float)):
            return int(round(float(val)))
        # Clean string like "800.649.036" or "800,649,036"
        s = str(val).strip().replace(".", "").replace(",", "")
        return int(round(float(s)))
    except:
        return 0


def parse_thqt_excel(file_source):
    """
    file_source: có thể là đường dẫn file (str) hoặc BytesIO / bytes.
    Trả về dict chứa đầy đủ thông tin bóc tách.
    """
    if isinstance(file_source, bytes):
        wb = openpyxl.load_workbook(BytesIO(file_source), data_only=True)
    elif isinstance(file_source, str):
        wb = openpyxl.load_workbook(file_source, data_only=True)
    elif hasattr(file_source, 'read'):
        wb = openpyxl.load_workbook(file_source, data_only=True)
    else:
        raise ValueError("Invalid file_source for parse_thqt_excel")

    sheet_names = wb.sheetnames
    target_sheet = "Sheet1" if "Sheet1" in sheet_names else sheet_names[0]
    ws = wb[target_sheet]

    ten_ct = ""
    ma_ct = ""

    # Quét tiêu đề từ dòng 1 đến dòng 10
    for r in range(1, 11):
        for c in range(1, 8):
            val = str(ws.cell(r, c).value or "")
            if "Tên công trình:" in val:
                ten_ct = val.split("Tên công trình:")[1].strip()
            elif "Mã Công Trình:" in val or "Mã công trình:" in val:
                m = re.split(r"Mã [cC]ông [tT]rình:", val)
                if len(m) > 1:
                    ma_ct = m[1].strip()

    # Bóc tách 44 dòng dữ liệu
    rows_data = []
    # Tìm dòng bắt đầu header
    start_row = 11
    for r in range(1, 15):
        val = str(ws.cell(r, 1).value or "").strip()
        if "Nội dung chi phí" in val:
            start_row = r + 1
            break

    row_map = {}
    for r in range(start_row, 46):
        name_val = ws.cell(r, 1).value
        if name_val is None or str(name_val).strip() == "":
            continue
        name = str(name_val).strip()
        dt = _clean_num(ws.cell(r, 2).value)
        ps_tang = _clean_num(ws.cell(r, 3).value)
        ps_giam = _clean_num(ws.cell(r, 4).value)
        tong_dt = _clean_num(ws.cell(r, 5).value)
        qt = _clean_num(ws.cell(r, 6).value)
        note = str(ws.cell(r, 7).value or "").strip()

        item = {
            "row_idx": r,
            "name": name,
            "dt": dt,
            "ps_tang": ps_tang,
            "ps_giam": ps_giam,
            "tong_dt": tong_dt,
            "qt": qt,
            "note": note,
        }
        rows_data.append(item)
        row_map[name.lower()] = item

    # Hàm tìm giá trị linh hoạt theo keyword hoặc theo số dòng
    def get_val(keyword, default_row_idx, col="qt"):
        for it in rows_data:
            if keyword.lower() in it["name"].lower():
                return it[col]
        # Fallback theo row_idx
        for it in rows_data:
            if it["row_idx"] == default_row_idx:
                return it[col]
        return 0

    sc_dt = get_val("Giá trị sửa chữa trước thuế", 27, "dt")
    sc_qt = get_val("Giá trị sửa chữa trước thuế", 27, "qt")

    tb_dt = get_val("Chi phí thiết bị, vật tư (trước thuế)", 19, "dt")
    tb_qt = get_val("Chi phí thiết bị, vật tư (trước thuế)", 19, "qt")

    khac_dt = get_val("Chi phí khác trước thuế", 36, "dt")
    khac_qt = get_val("Chi phí khác trước thuế", 36, "qt")

    truoc_thue_dt = get_val("Giá trị dự toán trước thuế", 40, "dt")
    truoc_thue_qt = get_val("Giá trị dự toán trước thuế", 40, "qt")
    if truoc_thue_qt == 0 and (sc_qt > 0 or tb_qt > 0 or khac_qt > 0):
        truoc_thue_qt = sc_qt + tb_qt + khac_qt
    if truoc_thue_dt == 0 and (sc_dt > 0 or tb_dt > 0 or khac_dt > 0):
        truoc_thue_dt = sc_dt + tb_dt + khac_dt

    thue_dt = get_val("2. Thuế GTGT", 41, "dt")
    thue_qt = get_val("2. Thuế GTGT", 41, "qt")

    tong_dt = get_val("Giá trị dự toán (không tính VTTH)", 39, "dt")
    tong_qt = get_val("Giá trị dự toán (không tính VTTH)", 39, "qt")
    if tong_qt == 0 and truoc_thue_qt > 0:
        tong_qt = truoc_thue_qt + thue_qt
    if tong_dt == 0 and truoc_thue_dt > 0:
        tong_dt = truoc_thue_dt + thue_dt

    vtth_dt = get_val("Giá trị VTTB thu hồi", 42, "dt")
    vtth_qt = get_val("Giá trị VTTB thu hồi", 42, "qt")

    scl_dt = get_val("Tổng chi phí sửa chữa lớn", 44, "dt")
    scl_qt = get_val("Tổng chi phí sửa chữa lớn", 44, "qt")
    if scl_qt == 0 and truoc_thue_qt > 0:
        scl_qt = truoc_thue_qt - vtth_qt
    if scl_dt == 0 and truoc_thue_dt > 0:
        scl_dt = truoc_thue_dt - vtth_dt

    # Bóc tách các mục chi phí KTCB khác cho Table 1 của Báo cáo Thẩm tra
    ktcb_items = [
        {"name": "Chi phí Lập khảo sát, Lập PAKT-DT", "def_row": 30},
        {"name": "2- Chi phí giám sát công trình", "def_row": 31},
        {"name": "Chi phí bảo hiểm công trình", "def_row": 32},
        {"name": "Chi phí thẩm tra thiết kế", "def_row": 33},
        {"name": "Chi phí thẩm tra dự toán", "def_row": 34},
        {"name": "Chi phí thẩm tra phê duyệt quyết toán", "def_row": 35},
        {"name": "Chi phí dự phòng", "def_row": 38},
    ]

    ktcb_rows = []
    tot_ktcb_dt = 0
    tot_ktcb_qt = 0
    for ki in ktcb_items:
        r_dt = get_val(ki["name"], ki["def_row"], "dt")
        r_qt = get_val(ki["name"], ki["def_row"], "qt")
        tot_ktcb_dt += r_dt
        tot_ktcb_qt += r_qt
        ktcb_rows.append({
            "name": ki["name"],
            "dt": r_dt,
            "qt": r_qt,
            "tt": r_qt,
            "diff": 0,
        })

    # Dòng tổng cộng KTCB
    # Nếu r_qt các mục trên bằng 0 nhưng khac_qt > 0 thì gán khac_qt
    if tot_ktcb_qt == 0 and khac_qt > 0:
        tot_ktcb_qt = khac_qt

    summary = {
        "sc_dt": sc_dt,
        "sc_qt": sc_qt,
        "tb_dt": tb_dt,
        "tb_qt": tb_qt,
        "khac_dt": khac_dt,
        "khac_qt": khac_qt,
        "truoc_thue_dt": truoc_thue_dt,
        "truoc_thue_qt": truoc_thue_qt,
        "thue_dt": thue_dt,
        "thue_qt": thue_qt,
        "tong_dt": tong_dt,
        "tong_qt": tong_qt,
        "vtth_dt": vtth_dt,
        "vtth_qt": vtth_qt,
        "scl_dt": scl_dt,
        "scl_qt": scl_qt,
        "bang_chu": doc_so_vn(tong_qt) if tong_qt > 0 else "",
        "tot_ktcb_dt": tot_ktcb_dt,
        "tot_ktcb_qt": tot_ktcb_qt,
    }

    return {
        "ten_ct": ten_ct,
        "ma_ct": ma_ct,
        "summary": summary,
        "ktcb_rows": ktcb_rows,
        "rows_data": rows_data,
    }


def save_thqt_cache(ma_ct, data, raw_bytes=None):
    """Lưu dữ liệu Bảng THQT vào thư mục data để dùng bền vững."""
    if not ma_ct:
        return
    ma_clean = re.sub(r'[^a-zA-Z0-9_\-]', '', str(ma_ct).strip())
    json_path = os.path.join(DATA_DIR, f"thqt_{ma_clean}.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    if raw_bytes:
        xlsx_path = os.path.join(DATA_DIR, f"thqt_{ma_clean}.xlsx")
        with open(xlsx_path, "wb") as f:
            f.write(raw_bytes)


def load_thqt_cache(ma_ct):
    """Tải dữ liệu Bảng THQT từ cache nếu có."""
    if not ma_ct:
        return None
    ma_clean = re.sub(r'[^a-zA-Z0-9_\-]', '', str(ma_ct).strip())
    json_path = os.path.join(DATA_DIR, f"thqt_{ma_clean}.json")
    if os.path.exists(json_path):
        try:
            with open(json_path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return None

    # Kiểm tra nếu trùng với mã trong HS_mau (ví dụ VTAD2607001)
    hs_mau_path = os.path.join(BASE_DIR, "HS_mau", "2. Bảng THQT KP SCL.xlsx")
    if os.path.exists(hs_mau_path):
        try:
            res = parse_thqt_excel(hs_mau_path)
            if res.get("ma_ct") == ma_ct:
                save_thqt_cache(ma_ct, res)
                return res
        except:
            pass

    return None
